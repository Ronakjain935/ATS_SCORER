import io
import os
from typing import List, Optional, Tuple

try:
    import magic
except (ImportError, Exception):
    magic = None

import pdfplumber
from docx import Document
import PyPDF2

from backend.utils.file_utils import (
    FileParsingError,
    TextExtractionError,
    FileUploadError,
    log_error,
    log_warning,
    log_info,
    with_fallback,
)

from backend.core.config import (
    MAX_FILE_SIZE_BYTES,
    MAX_FILE_SIZE_MB,
    SUPPORTED_MIME_TYPES,
)


class FileValidationError(FileUploadError):
    """Raised when an uploaded file fails validation (size, type, empty, corrupted)."""
    pass


def _detect_file_type(file_data: bytes, filename: str) -> Optional[str]:
    """
    Detect file format (pdf, docx, doc) using a combination of:
    1. python-magic MIME detection (if available)
    2. Magic bytes / header sniffing
    3. File extension from filename
    """
    ext = os.path.splitext(filename or '')[1].lower().lstrip('.')

    # 1. Try python-magic if available
    detected_mime = None
    if magic is not None:
        try:
            detected_mime = magic.from_buffer(file_data, mime=True)
        except Exception:
            detected_mime = None

    if detected_mime and detected_mime in SUPPORTED_MIME_TYPES:
        return SUPPORTED_MIME_TYPES[detected_mime]

    # 2. Header magic bytes sniffing
    # PDF starts with %PDF
    if file_data.startswith(b'%PDF'):
        return 'pdf'

    # DOCX is a ZIP container starting with PK\x03\x04
    if file_data.startswith(b'PK\x03\x04'):
        if ext == 'docx' or not ext or detected_mime in ('application/zip', 'application/x-zip', 'application/octet-stream'):
            return 'docx'

    # Legacy DOC starts with OLE Compound Document signature \xd0\xcf\x11\xe0
    if file_data.startswith(b'\xd0\xcf\x11\xe0'):
        return 'doc'

    # 3. Fallback to extension check if matching supported types
    if ext in ('pdf', 'docx', 'doc'):
        return ext

    return None


def validate_file(file_data: bytes, filename: str) -> Tuple[bool, str, Optional[str]]:
    file_size_bytes = len(file_data)

    if file_size_bytes == 0:
        return False, 'Uploaded file is empty. Please check the file and try again.', None

    if file_size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = file_size_bytes / (1024 * 1024)
        return False, (
            f'File size ({size_mb:.2f} MB) exceeds the maximum allowed size of {MAX_FILE_SIZE_MB} MB. '
            'Please upload a smaller file or compress your resume.'
        ), None

    file_type = _detect_file_type(file_data, filename)
    if not file_type:
        return False, (
            f"Unsupported file format for '{filename}'. "
            "Please upload a valid PDF (.pdf) or Word document (.docx)."
        ), None

    if file_type == 'doc':
        return False, (
            'Legacy .doc format is not supported. '
            'Please convert your resume to .docx or .pdf and try again.'
        ), None

    return True, '', file_type


def _extract_pdf_hyperlinks(file_data: bytes) -> str:
    urls = []
    seen = set()
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_data))
        for page in reader.pages:
            annots = page.get('/Annots')
            if not annots:
                continue
            for annot_ref in annots:
                try:
                    annot = annot_ref.get_object() if hasattr(annot_ref, 'get_object') else annot_ref
                    if not isinstance(annot, dict) or annot.get('/Subtype') != '/Link':
                        continue
                    action = annot.get('/A', {})
                    if hasattr(action, 'get_object'):
                        action = action.get_object()
                    uri = action.get('/URI', '') if isinstance(action, dict) else ''
                    if uri and isinstance(uri, (str, bytes)):
                        if isinstance(uri, bytes):
                            uri = uri.decode('utf-8', errors='ignore')
                        uri = uri.strip()
                        if uri.startswith(('http://', 'https://')) and uri not in seen:
                            seen.add(uri)
                            urls.append(uri)
                except Exception:
                    pass
    except Exception:
        pass
    return '\n'.join(urls)


def _extract_pdf_with_pdfplumber(file_data: bytes) -> str:
    text = ''
    with pdfplumber.open(io.BytesIO(file_data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + '\n'

    if not text.strip():
        raise TextExtractionError(
            'pdfplumber extracted no text',
            user_message='No text could be extracted from the PDF.'
        )

    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()


def _extract_pdf_with_pypdf2(file_data: bytes) -> str:
    text = ''
    pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_data))
    for page in pdf_reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + '\n'

    if not text.strip():
        raise TextExtractionError(
            'PyPDF2 extracted no text',
            user_message='No text could be extracted from the PDF.'
        )

    hyperlinks = _extract_pdf_hyperlinks(file_data)
    if hyperlinks:
        text = text.strip() + '\n' + hyperlinks

    return text.strip()


def extract_text_from_pdf(file_data: bytes) -> str:
    try:
        result, used_fallback = with_fallback(
            _extract_pdf_with_pdfplumber,
            _extract_pdf_with_pypdf2,
            file_data,
            log_fallback=True
        )

        if used_fallback:
            log_info('PDF extraction succeeded using PyPDF2 fallback', context='resume_parser')
        return result

    except Exception as e:
        log_error(e, context='extract_text_from_pdf')
        raise FileParsingError(
            'Failed to extract text from PDF using both pdfplumber and PyPDF2. '
            'The PDF may be corrupted, password-protected, or contain only scanned images. '
            'Please ensure it contains selectable text.',
            user_message='Could not read text from this PDF. Please ensure it is not scanned, password-protected, or empty.'
        ) from e


def extract_text_from_docx(file_data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_data))
        text_parts: List[str] = []

        for paragraph in doc.paragraphs:
            p_text = paragraph.text.strip()
            if p_text:
                text_parts.append(p_text)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate adjacent duplicate cells (common with merged cells in tables)
                    unique_cells: List[str] = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    text_parts.append(' | '.join(unique_cells))

        text = '\n'.join(text_parts)

        # Extract embedded hyperlinks
        urls: List[str] = []
        seen = set()
        try:
            for rel in doc.part.rels.values():
                reltype = getattr(rel, 'reltype', '') or ''
                if 'hyperlink' in reltype.lower():
                    target = getattr(rel, 'target_ref', None) or getattr(rel, '_target', None)
                    if isinstance(target, str) and target.startswith(('http://', 'https://')):
                        target = target.strip()
                        if target not in seen:
                            seen.add(target)
                            urls.append(target)
        except Exception:
            pass

        if urls:
            text = (text + '\n' + '\n'.join(urls)).strip()

        if not text.strip():
            raise FileParsingError(
                'No text could be extracted from the document. The document may be empty or corrupted.',
                user_message='No text could be extracted from the Word document.'
            )

        log_info(f'Extracted {len(text)} chars from DOCX', context='resume_parser')
        return text.strip()

    except FileParsingError:
        raise
    except Exception as e:
        log_error(e, context='extract_text_from_docx')
        raise FileParsingError(
            'Failed to extract text from DOCX. The document may be corrupted or in an unsupported format.',
            user_message='Failed to extract text from DOCX. Please ensure it is a valid .docx file or convert to PDF.'
        ) from e


def extract_text_from_doc(file_data: bytes) -> str:
    raise FileParsingError(
        'Legacy .doc format is not supported. '
        'Please convert your document to .docx or .pdf and try again.',
        user_message='Legacy .doc format is not supported. Please convert to .docx or .pdf.'
    )


def extract_text(file_data: bytes, file_type: str) -> str:
    normalized_type = (file_type or '').strip().lower().lstrip('.')
    if normalized_type == 'pdf':
        return extract_text_from_pdf(file_data)
    elif normalized_type == 'docx':
        return extract_text_from_docx(file_data)
    elif normalized_type == 'doc':
        return extract_text_from_doc(file_data)
    else:
        raise FileValidationError(
            f'Invalid file type: {file_type}. Supported types are: pdf and docx.',
            user_message=f"Unsupported file format '{file_type}'. Please upload a PDF or DOCX file."
        )


def parse_resume_file(file_data: bytes, filename: str) -> Tuple[str, dict]:
    log_info(f'Parsing file: {filename}', context='parse_resume_file')

    # Phase 01: Validate file
    try:
        is_valid, error_msg, file_type = validate_file(file_data, filename)
        if not is_valid or file_type is None:
            err = error_msg or f"Unsupported file format for '{filename}'."
            log_warning(f'Validation failed for file {filename}: {err}', context='parse_resume_file')
            raise FileValidationError(err, user_message=err)
    except FileValidationError:
        raise
    except Exception as e:
        log_error(e, context='parse_resume_file_validation')
        raise FileValidationError(
            'Could not validate the uploaded file. Please ensure it is a valid PDF or DOCX.',
            user_message='Could not validate the uploaded file. Please ensure it is a valid PDF or DOCX.'
        ) from e

    # Phase 02: Extraction of file text
    try:
        text = extract_text(file_data, file_type)
        log_info(f'Extracted {len(text)} chars from {filename}', context='parse_resume_file')
    except (FileParsingError, FileValidationError):
        raise
    except Exception as e:
        log_error(e, context='parse_resume_file_extraction')
        raise FileParsingError(
            'An unexpected error occurred while processing the file. '
            'Please try again or contact support if the problem persists.',
            user_message='An unexpected error occurred while extracting text from the file.'
        ) from e

    metadata = {
        'filename': filename,
        'file_type': file_type,
        'file_size_bytes': len(file_data),
        'text_length': len(text),
        'success': True,
    }
    return text, metadata